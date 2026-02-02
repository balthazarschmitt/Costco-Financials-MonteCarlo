import openpyxl as pyxl
import numpy as np


# Load the workbook and select the sheet
wb = pyxl.load_workbook('Costco Financial Model.xlsx', data_only=True)
ws = wb['3_Statement_Model']

revenue_0 = ws['H6'].value
tax_rate = ws['I23'].value
net_debt = wb['DCF']['N80'].value
shares = wb['DCF']['N82'].value

N = 20000

growth_1_5 = np.random.normal(0.06, 0.015, N)
growth_6_10 = np.random.normal(0.04, 0.01, N)
ebit_margin = np.random.normal(0.037, 0.003, N)

wacc = np.random.normal(0.0734, 0.005, N)
tgr = np.random.triangular(0.025, 0.03, 0.035, N)

wacc = np.clip(wacc, 0.05, 0.12)
tgr = np.clip(tgr, 0.0, 0.04)

def dcf_value(revenue_0, g1, g2, margin, tax_rate, wacc, tgr, net_debt):
    revenues = []
    revenue = revenue_0
    
    for t in range(1, 11):
        g = g1 if t <= 5 else g2
        revenue *= (1 + g)
        revenues.append(revenue)
        
    ebit = np.array(revenues) * margin
    nopat = ebit * (1 - tax_rate)
    
    fcff = nopat
    
    discount_factors = np.array([(1 + wacc)**t for t in range(1, 11)])
    pv_fcff = np.sum(fcff / discount_factors)
    
    terminal_value = fcff[-1] * (1 + tgr) / (wacc - tgr)
    pv_terminal = terminal_value / ((1 + wacc) ** 10)
    
    enterprise_value = pv_fcff + pv_terminal
    equity_value = enterprise_value - net_debt
    
    return equity_value

equity_values = np.array([
    dcf_value(
        revenue_0,
        growth_1_5[i],
        growth_6_10[i],
        ebit_margin[i],
        tax_rate,
        wacc[i],
        tgr[i],
        net_debt
    )
    for i in range(N)
])

value_per_share = equity_values / shares

np.percentile(value_per_share, [5, 25, 50, 75, 95])
print("10th Percentile: ", np.percentile(value_per_share, 10))
print("25th Percentile: ", np.percentile(value_per_share, 25))
print("50th Percentile: ", np.percentile(value_per_share, 50))
print("75th Percentile: ", np.percentile(value_per_share, 75))
print("90th Percentile: ", np.percentile(value_per_share, 90))
